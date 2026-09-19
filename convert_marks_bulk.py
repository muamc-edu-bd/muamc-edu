"""
convert_marks_bulk.py
=====================
Converts "Science-All (1).xlsx" and "Business_All (1).xlsx" marksheets
into the Excel format required by the Bulk Marks Import feature
on marks-entry.html.

Expected output format columns:
  Roll | Exam Type | Session | Optional Subject Code | {subject_code_1} | ...

Each subject column contains marks as: CQ/MCQ or CQ/MCQ/Practical
Absent students ('Ab') are skipped entirely (cells left blank).
N/A subjects (not applicable to student) are also left blank.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Configuration ────────────────────────────────────────────────────────────
EXAM_TYPE = "Annual"
SESSION   = "2025-2026"
INPUT_DIR = "Excel"
OUTPUT_DIR = "Excel"


# ── Utility helpers ──────────────────────────────────────────────────────────

def is_valid_mark(val):
    """Check if a value is a valid numeric mark (not None, N/A, Ab, or empty)."""
    if val is None:
        return False
    s = str(val).strip()
    return s not in ('', 'N/A', 'n/a', 'Ab', 'ab', 'AB', '-')


def is_absent(val):
    """Check if a value represents an absent mark."""
    if val is None:
        return False
    s = str(val).strip()
    return s.lower() == 'ab'


def safe_int(val):
    """Convert a value to integer, returning 0 for invalid values."""
    if val is None:
        return 0
    s = str(val).strip()
    if s in ('', 'N/A', 'n/a', 'Ab', 'ab', 'AB', '-'):
        return 0
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return 0


def format_marks(cq, mcq, prac=None):
    """Format marks as CQ/MCQ or CQ/MCQ/Practical string."""
    if prac is not None:
        return f"{safe_int(cq)}/{safe_int(mcq)}/{safe_int(prac)}"
    return f"{safe_int(cq)}/{safe_int(mcq)}"


def is_numeric_roll(roll_str):
    """Return True if the roll string represents a number."""
    try:
        int(str(roll_str).strip())
        return True
    except (ValueError, TypeError):
        return False


def subject_has_any_absent(vals):
    """Check if any component value in a subject is 'Ab' (absent)."""
    return any(is_absent(v) for v in vals)


def subject_all_na_or_empty(vals):
    """Check if ALL component values are N/A or empty (not applicable)."""
    return all(not is_valid_mark(v) for v in vals)


# ── Science converter ────────────────────────────────────────────────────────

def convert_science():
    """
    Convert Science-All (1).xlsx → Science_Bulk_Marks_Ready.xlsx

    Source column layout (0-indexed):
      0: Roll  |  1: Name  |  2: Registration No.
      3-4:   Bangla 1st Paper    (CQ, MCQ)           → code 101
      5-6:   English 1st Paper   (CQ, MCQ)           → code 107
      7-9:   ICT                 (CQ, MCQ, Practical) → code 275
      10-12: Physics 1st Paper   (CQ, MCQ, Practical) → code 174
      13-15: Chemistry 1st Paper (CQ, MCQ, Practical) → code 176
      16-18: Biology 1st Paper   (CQ, MCQ, Practical) → code 178
      19-21: Higher Math 1st     (CQ, MCQ, Practical) → code 265
      22-24: Total / GPA / Grade (ignored)
    """
    wb = openpyxl.load_workbook(f"{INPUT_DIR}/Science-All (1).xlsx", data_only=True)
    ws = wb.active

    # (subject_code, cq_col, mcq_col, prac_col_or_None)
    SUBJECT_MAP = [
        ('101', 3,  4,  None),   # Bangla 1st Paper
        ('107', 5,  6,  None),   # English 1st Paper
        ('275', 7,  8,  9),      # ICT
        ('174', 10, 11, 12),     # Physics 1st Paper
        ('176', 13, 14, 15),     # Chemistry 1st Paper
        ('178', 16, 17, 18),     # Biology 1st Paper
        ('265', 19, 20, 21),     # Higher Math 1st Paper
    ]

    CODES = [s[0] for s in SUBJECT_MAP]
    headers = ['Roll', 'Exam Type', 'Session', 'Optional Subject Code'] + CODES

    rows_out = []
    skipped = 0

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue  # Skip the 2 header rows

        roll_raw = str(row[0] or '').strip()
        if not is_numeric_roll(roll_raw):
            skipped += 1
            continue

        roll = str(int(roll_raw))   # Normalize: "01" → "1"

        # ── Determine optional subject ──
        # Most students have both Biology and Higher Math marks.
        # If only one is present, we set the optional code.
        # If both are present, we leave it blank → system uses existing assignment.
        bio_vals  = [row[16], row[17], row[18]]
        hm_vals   = [row[19], row[20], row[21]]
        has_bio   = not subject_all_na_or_empty(bio_vals) and not subject_has_any_absent(bio_vals)
        has_hm    = not subject_all_na_or_empty(hm_vals)  and not subject_has_any_absent(hm_vals)

        if has_bio and not has_hm:
            optional_code = '178/179'
        elif has_hm and not has_bio:
            optional_code = '265/266'
        else:
            optional_code = ''   # Both present or neither → use existing assignment

        out_row = {
            'Roll': roll,
            'Exam Type': EXAM_TYPE,
            'Session': SESSION,
            'Optional Subject Code': optional_code,
        }

        for code, cq_col, mcq_col, prac_col in SUBJECT_MAP:
            vals = [row[cq_col], row[mcq_col]]
            if prac_col is not None:
                vals.append(row[prac_col])

            # Skip if subject is not applicable (all N/A or empty)
            if subject_all_na_or_empty(vals):
                continue

            # Skip if student was absent (user requested: leave blank)
            if subject_has_any_absent(vals):
                continue

            out_row[code] = format_marks(
                row[cq_col], row[mcq_col],
                row[prac_col] if prac_col is not None else None
            )

        rows_out.append(out_row)

    output_path = f"{OUTPUT_DIR}/Science_Bulk_Marks_Ready.xlsx"
    write_output(output_path, "Science Bulk Marks", headers, rows_out)

    print(f"  Source rows scanned : {ws.max_row - 2}")
    print(f"  Skipped (non-roll)  : {skipped}")
    print(f"  Students written    : {len(rows_out)}")
    print(f"  Output file         : {output_path}")

    return len(rows_out)


# ── Business converter ───────────────────────────────────────────────────────

def convert_business():
    """
    Convert Business_All (1).xlsx → Business_Bulk_Marks_Ready.xlsx

    Source column layout (0-indexed):
      0: Roll  |  1: Name  |  2: Registration No.
      3-4:   Bangla 1st Paper              (CQ, MCQ)           → code 101
      5-6:   English 1st Paper             (CQ, MCQ)           → code 107
      7-9:   ICT                           (CQ, MCQ, Practical) → code 275
      10-11: Accounting 1st Paper          (CQ, MCQ)           → code 253
      12-13: Business Org. & Mgmt 1st      (CQ, MCQ)           → code 277
      14-15: Finance, Banking & Insurance  (CQ, MCQ)           → code 292
      16-17: Economics 1st Paper           (CQ, MCQ)           → code 109  [Optional]
      18-20: Home Science 1st Paper        (CQ, MCQ, Practical) → code 273 [Optional]
      21-23: Total / GPA / Grade (ignored)
    """
    wb = openpyxl.load_workbook(f"{INPUT_DIR}/Business_All (1).xlsx", data_only=True)
    ws = wb.active

    SUBJECT_MAP = [
        ('101', 3,  4,  None),   # Bangla 1st Paper
        ('107', 5,  6,  None),   # English 1st Paper
        ('275', 7,  8,  9),      # ICT
        ('253', 10, 11, None),   # Accounting 1st Paper
        ('277', 12, 13, None),   # Business Org. & Mgmt 1st Paper
        ('292', 14, 15, None),   # Finance, Banking & Insurance 1st Paper
        ('109', 16, 17, None),   # Economics 1st Paper [Optional]
        ('273', 18, 19, 20),     # Home Science 1st Paper [Optional]
    ]

    CODES = [s[0] for s in SUBJECT_MAP]
    headers = ['Roll', 'Exam Type', 'Session', 'Optional Subject Code'] + CODES

    rows_out = []
    skipped = 0

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue  # Skip the 2 header rows

        roll_raw = str(row[0] or '').strip()
        if not is_numeric_roll(roll_raw):
            skipped += 1
            continue

        roll = str(int(roll_raw))

        # ── Determine optional subject ──
        econ_vals = [row[16], row[17]]
        hs_vals   = [row[18], row[19], row[20]]

        has_econ = not subject_all_na_or_empty(econ_vals) and not subject_has_any_absent(econ_vals)
        has_hs   = not subject_all_na_or_empty(hs_vals)   and not subject_has_any_absent(hs_vals)

        if has_econ:
            optional_code = '109/110'   # busEcon
        elif has_hs:
            optional_code = '273/274'   # busHome
        else:
            optional_code = ''          # Absent on optional → leave blank

        out_row = {
            'Roll': roll,
            'Exam Type': EXAM_TYPE,
            'Session': SESSION,
            'Optional Subject Code': optional_code,
        }

        for code, cq_col, mcq_col, prac_col in SUBJECT_MAP:
            vals = [row[cq_col], row[mcq_col]]
            if prac_col is not None:
                vals.append(row[prac_col])

            if subject_all_na_or_empty(vals):
                continue

            if subject_has_any_absent(vals):
                continue

            out_row[code] = format_marks(
                row[cq_col], row[mcq_col],
                row[prac_col] if prac_col is not None else None
            )

        rows_out.append(out_row)

    output_path = f"{OUTPUT_DIR}/Business_Bulk_Marks_Ready.xlsx"
    write_output(output_path, "Business Bulk Marks", headers, rows_out)

    print(f"  Source rows scanned : {ws.max_row - 2}")
    print(f"  Skipped (non-roll)  : {skipped}")
    print(f"  Students written    : {len(rows_out)}")
    print(f"  Output file         : {output_path}")

    return len(rows_out)


# ── Output writer ────────────────────────────────────────────────────────────

def write_output(filename, sheet_title, headers, rows):
    """Write a styled bulk-upload Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Styles
    hdr_font  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    hdr_fill  = PatternFill("solid", fgColor="1A3C5E")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    dat_align = Alignment(horizontal="center", vertical="center")
    dat_font  = Font(name="Calibri", size=10)
    border    = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    alt_fill  = PatternFill("solid", fgColor="EBF2FB")

    # ── Header row ──
    col_widths = {
        'Roll': 8, 'Exam Type': 14, 'Session': 14, 'Optional Subject Code': 24,
    }
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = col_widths.get(h, 14)

    ws.row_dimensions[1].height = 22

    # ── Data rows ──
    for row_num, row_data in enumerate(rows, start=2):
        fill = alt_fill if row_num % 2 == 0 else None
        for col, h in enumerate(headers, start=1):
            val = row_data.get(h, '')
            cell = ws.cell(row=row_num, column=col, value=val if val else '')
            cell.alignment = dat_align
            cell.border = border
            cell.font = dat_font
            if fill:
                cell.fill = fill

    # Freeze header & auto-filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    wb.save(filename)


# ── Main ─────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print("=" * 60)
    print("  Marksheet -> Bulk Upload Converter")
    print("=" * 60)
    print(f"  Exam Type : {EXAM_TYPE}")
    print(f"  Session   : {SESSION}")
    print(f"  Absent    : Skipped (left blank)")
    print()

    print("[1/2] Converting Science marksheet...")
    sci_count = convert_science()
    print()

    print("[2/2] Converting Business marksheet...")
    bus_count = convert_business()
    print()

    print("=" * 60)
    print(f"  Total students processed: {sci_count + bus_count}")
    print(f"  Output files in: {OUTPUT_DIR}/")
    print("  -> Science_Bulk_Marks_Ready.xlsx")
    print("  -> Business_Bulk_Marks_Ready.xlsx")
    print()
    print("  Upload these via: Marks Entry -> Marks Bulk Input")
    print("=" * 60)
